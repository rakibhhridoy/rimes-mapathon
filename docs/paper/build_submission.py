"""
Build the Journal of Flood Risk Management submission package from paper.tex.

The journal wants the main text in Word, figures as separate files with their
legends listed in the text, tables with titles, and APA references. This
script writes, into docs/paper/submission/:

    manuscript.docx        title, abstract, keywords, main text, declarations,
                           data availability, references (APA 7), tables,
                           figure legends
    title_page.docx        title, running title, author, affiliations, ORCID,
                           correspondence, acknowledgments
    figures/FigureN.pdf    each figure as a vector PDF, and FigureN.png at 300 dpi
    supporting_information.pdf
    cover_letter.docx      from cover_letter.md, with a PDF copy
    upload/                only the files to upload, numbered in upload order

Every number comes from docs/numbers.tex and every figure, table and section
number from paper.aux, so the Word file matches the compiled PDF. Compile
paper.tex (and supporting_information.tex) first.

    python docs/paper/build_submission.py
"""

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
DOCS = HERE.parent
OUT = HERE / "submission"
CSL = OUT / "apa.csl"
BIB = HERE / "references.bib"


def macros() -> dict:
    """\\newcommand{\\Name}{value} from numbers.tex."""
    out = {}
    for line in (DOCS / "numbers.tex").read_text().splitlines():
        m = re.match(r"\\newcommand\{\\([A-Za-z]+)\}\{(.*)\}$", line)
        if m:
            out[m.group(1)] = m.group(2)
    return out


def labels() -> dict:
    """label -> number, as LaTeX numbered it in the last compile."""
    out = {}
    for line in (HERE / "paper.aux").read_text().splitlines():
        m = re.match(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", line)
        if m:
            out[m.group(1)] = m.group(2)
    return out


def expand(tex: str, nums: dict, labs: dict) -> str:
    def macro(m):
        name = m.group(1)
        if name not in nums:
            return m.group(0)
        value = nums[name]
        if "todo" in value:
            raise ValueError(f"\\{name} has no value in numbers.tex")
        return value
    tex = re.sub(r"\\([A-Za-z]+)\{\}", macro, tex)
    tex = re.sub(r"\\ref\{([^}]*)\}", lambda m: labs[m.group(1)], tex)
    tex = tex.replace(r"\fileref", r"\texttt")
    return tex


def simplify_table(env: str) -> str:
    """pandoc reads plain tabular reliably; tabularx column specs it does not."""
    env = env.replace("table*", "table")
    m = re.search(r"\\begin\{tabularx\}\{[^}]*\}\{", env)
    if m:
        depth, i = 1, m.end()
        while depth:
            depth += {"{": 1, "}": -1}.get(env[i], 0)
            i += 1
        first_row = env[i:].split(r"\\")[0]
        ncols = first_row.count("&") + 1
        env = env[:m.start()] + r"\begin{tabular}{" + "l" * ncols + "}" + env[i:]
        env = env.replace(r"\end{tabularx}", r"\end{tabular}")
    return env


def to_markdown(tex: str) -> str:
    with tempfile.NamedTemporaryFile("w", suffix=".tex", delete=False) as f:
        f.write(tex)
        path = f.name
    out = subprocess.run(["pandoc", "-f", "latex", "-t", "markdown", "--wrap=none", path],
                         capture_output=True, text=True, check=True).stdout
    Path(path).unlink()
    return out


def reference_doc(path: Path) -> Path:
    """pandoc's default Word styles with Times New Roman 12 pt throughout and
    black headings; the defaults use theme fonts and blue headings."""
    import zipfile

    from docx import Document
    from docx.shared import Pt

    default = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                             capture_output=True, check=True).stdout
    raw = path.with_suffix(".raw.docx")
    raw.write_bytes(default)
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/styles.xml":
                xml = data.decode()
                xml = re.sub(r"<w:rFonts\b[^>]*/>", '<w:rFonts w:ascii="Times New Roman" '
                             'w:hAnsi="Times New Roman" w:eastAsia="Times New Roman" '
                             'w:cs="Times New Roman" />', xml)
                xml = re.sub(r"<w:color\b[^>]*/>", '<w:color w:val="000000" />', xml)
                data = xml.encode()
            dst.writestr(item, data)
    raw.unlink()
    doc = Document(str(path))
    for style in doc.styles:
        if style.type != 1:
            continue
        font = style.font
        font.name = "Times New Roman"
        if style.name.startswith("Heading") or style.name in ("Title", "Subtitle"):
            font.size = Pt(14 if style.name == "Title" else 12)
            font.bold = style.name in ("Title", "Heading 1", "Heading 2")
            font.italic = style.name in ("Heading 2", "Heading 3")
        elif font.size and font.size > Pt(12):
            font.size = Pt(12)
    doc.save(str(path))
    return path


def justify(path: Path) -> None:
    """Justify running text; lines split by manual breaks (address, date)
    stay left-aligned, as Word would stretch them."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document(str(path))
    for p in doc.paragraphs:
        if not p._p.xpath(".//w:br | .//w:numPr") and len(p.text.split()) > 12:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.save(str(path))


def for_review(path: Path) -> None:
    """Double spacing, continuous line numbers and page numbers, so that
    reviewers can point to an exact line. The journal does not ask for them,
    but they cost nothing and are what reviewers expect."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(str(path))
    for style in doc.styles:
        if style.type == 1 and style.name in ("Normal", "Body Text", "First Paragraph",
                                              "Compact", "Bibliography"):
            style.paragraph_format.line_spacing = 2.0
    for section in doc.sections:
        numbering = OxmlElement("w:lnNumType")
        numbering.set(qn("w:countBy"), "1")
        numbering.set(qn("w:restart"), "continuous")
        section._sectPr.append(numbering)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer._p.get_or_add_pPr().append(OxmlElement("w:suppressLineNumbers"))
        run = footer.add_run()
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                field = OxmlElement("w:fldChar")
                field.set(qn("w:fldCharType"), kind)
            else:
                field = OxmlElement("w:instrText")
                field.set(qn("xml:space"), "preserve")
                field.text = text
            run._r.append(field)
    doc.save(str(path))


def turnitin(main_tex: str, nums: dict, labs: dict, style: list) -> None:
    """Main text and Supporting Information in one Word file for a similarity
    check: no title, author, affiliation, figures, legends, declarations or
    reference list. Citations stay in the text."""
    si = (HERE / "supporting_information.tex").read_text()
    si = si[si.index(r"\maketitle") + len(r"\maketitle"):si.index(r"\bibliographystyle")]
    si_labs = dict(labs)
    for line in (HERE / "supporting_information.aux").read_text().splitlines():
        m = re.match(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", line)
        if m:
            si_labs[m.group(1)] = m.group(2)
    si = re.sub(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}\n*", "", si, flags=re.S)

    def table(m):
        t = simplify_table(m.group(0))
        n = si_labs[re.search(r"\\label\{([^}]*)\}", t).group(1)]
        return re.sub(r"\\caption\{", f"\\\\caption{{Table {n}. ", t, count=1)
    si = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", table, si, flags=re.S)
    # Supporting Information sections keep their S numbers, unnumbered by pandoc.
    count = iter(range(1, 100))
    si = re.sub(r"\\section\{", lambda m: f"\\section*{{S{next(count)} ", si)
    md = [to_markdown(expand(main_tex, nums, labs)),
          "# Supporting Information {.unnumbered}\n",
          to_markdown(expand(si, nums, si_labs))]
    path = OUT / "turnitin.md"
    path.write_text("\n".join(md))
    subprocess.run(["pandoc", str(path), "-o", str(OUT / "turnitin.docx"),
                    "--citeproc", f"--bibliography={BIB}", f"--csl={CSL}",
                    "--number-sections", "--metadata", "link-citations=false",
                    "--metadata", "suppress-bibliography=true"] + style, check=True)


def main():
    nums, labs = macros(), labels()
    tex = (HERE / "paper.tex").read_text()
    OUT.mkdir(exist_ok=True)
    (OUT / "figures").mkdir(exist_ok=True)
    work = Path(tempfile.mkdtemp())
    style = ["--reference-doc", str(reference_doc(work / "reference.docx"))]

    title = re.search(r"\{\\LARGE\\bfseries (.*?)\\par\}", tex).group(1)
    running = re.search(r"Running title: (.*?)\\par", tex).group(1)
    abstract = expand((HERE / "abstract.tex").read_text().strip(), nums, labs)
    keywords = re.search(r"\\textbf\{Keywords:\} (.*)", tex).group(1).strip()

    body = tex[tex.index(r"\section{Introduction}"):tex.index(r"\section*{Acknowledgements}")]
    acknowledgments = tex[tex.index(r"\section*{Acknowledgements}") + len(r"\section*{Acknowledgements}"):
                          tex.index(r"\bibliographystyle")].strip()

    # Figures leave the text: legends go to a list, images to separate files.
    legends, figure_no = [], 0
    for m in re.finditer(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}", body, flags=re.S):
        fig = m.group(0)
        figure_no += 1
        n = labs[re.search(r"\\label\{([^}]*)\}", fig).group(1)]
        assert int(n) == figure_no, f"figure order {figure_no} != label {n}"
        src = DOCS / re.search(r"\\includegraphics(?:\[[^]]*\])?\{([^}]*)\}", fig).group(1)
        shutil.copy(src, OUT / "figures" / f"Figure{n}.pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "300", "-singlefile", str(src),
                        str(OUT / "figures" / f"Figure{n}")], check=True)
        caption = re.search(r"\\caption\{(.*)\}\s*\\label", fig, flags=re.S).group(1)
        legends.append(f"\\textbf{{Figure {n}.}} {caption}\n\n")
    body = re.sub(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}\n*", "", body, flags=re.S)

    # Tables move to the end, each titled with its number.
    tables = []
    for m in re.finditer(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", body, flags=re.S):
        t = simplify_table(m.group(0))
        n = labs[re.search(r"\\label\{([^}]*)\}", t).group(1)]
        t = re.sub(r"\\caption\{", f"\\\\caption{{Table {n}. ", t, count=1)
        tables.append((int(n), t))
    body = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}\n*", "", body, flags=re.S)
    tables.sort()

    main_tex = (f"\\section*{{Abstract}}\n\n{abstract}\n\n"
                f"\\textbf{{Keywords:}} {keywords}\n\n"
                + body)
    md = [to_markdown(expand(main_tex, nums, labs)),
          "# References {.unnumbered}\n\n::: {#refs}\n:::\n",
          "# Tables {.unnumbered}\n",
          to_markdown(expand("\n\n".join(t for _, t in tables), nums, labs)),
          "# Figure legends {.unnumbered}\n",
          to_markdown(expand("".join(legends), nums, labs))]
    # The similarity check reads the text and tables only, without declarations.
    turnitin(main_tex[:main_tex.index(r"\section*{Data availability statement}")]
             + "\n\n" + "\n\n".join(t for _, t in tables), nums, labs, style)
    manuscript_md = OUT / "manuscript.md"
    manuscript_md.write_text("\n".join(md))
    subprocess.run(["pandoc", str(manuscript_md), "-o", str(OUT / "manuscript.docx"),
                    "--citeproc", f"--bibliography={BIB}", f"--csl={CSL}",
                    "--number-sections", "--metadata", "link-citations=false",
                    "--metadata", f"title={title}"] + style,
                   check=True)
    for_review(OUT / "manuscript.docx")

    title_md = (f"**Running title:** {running}\n\n"
                "**Author:** Md Rakib Hasan^1,2^\n\n"
                "^1^ Department of Soil, Water and Environment, University of Dhaka, Dhaka-1000, Bangladesh\n\n"
                "^2^ Fermium Systems, Dhaka-1207, Bangladesh\n\n"
                "**ORCID:** 0009-0002-4007-7590\n\n"
                "**Correspondence:** Md Rakib Hasan, rakibhhridoy@fermium.systems\n\n"
                "## Acknowledgments\n\n"
                + to_markdown(expand(acknowledgments, nums, labs)))
    title_page = OUT / "title_page.md"
    title_page.write_text(title_md)
    subprocess.run(["pandoc", str(title_page), "-o", str(OUT / "title_page.docx"),
                    "--citeproc", f"--bibliography={BIB}", f"--csl={CSL}",
                    "--metadata", "suppress-bibliography=true",
                    "--metadata", f"title={title}"] + style, check=True)

    subprocess.run(["pandoc", str(HERE / "cover_letter.md"), "-o", str(OUT / "cover_letter.docx")]
                   + style, check=True)
    justify(OUT / "cover_letter.docx")
    # A PDF copy as well, for upload forms that take one; Times carries the Greek letters.
    subprocess.run(["pandoc", str(HERE / "cover_letter.md"), "-o", str(OUT / "cover_letter.pdf"),
                    "--pdf-engine=xelatex", "-V", "mainfont=Times New Roman", "-V", "fontsize=11pt",
                    "-V", "geometry:top=1.5cm,bottom=1.5cm,left=2cm,right=2cm",
                    "-V", "pagestyle=empty", "-V", "header-includes=\\usepackage{microtype}\\usepackage{xurl}", "-V", "linestretch=0.96"], check=True)

    si = HERE / "supporting_information.pdf"
    if si.exists():
        shutil.copy(si, OUT / "supporting_information.pdf")

    # Only the files the journal's form takes, numbered in upload order.
    upload = OUT / "upload"
    shutil.rmtree(upload, ignore_errors=True)
    upload.mkdir()
    files = (["cover_letter.pdf", "title_page.docx", "manuscript.docx"]
             + [f"figures/Figure{i}.pdf" for i in range(1, figure_no + 1)]
             + ["supporting_information.pdf"])
    for i, name in enumerate(files, 1):
        if (OUT / name).exists():
            shutil.copy(OUT / name, upload / f"{i:02d}_{Path(name).name}")
    print(f"wrote {OUT}: manuscript.docx, title_page.docx, {figure_no} figures, "
          f"{len(tables)} tables")


if __name__ == "__main__":
    main()
