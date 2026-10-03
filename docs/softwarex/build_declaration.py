"""
Write the SoftwareX declaration of interest as a Word file, in the wording of
Elsevier's standard declaration form. The disclosed interest repeats the
manuscript's "Declaration of competing interest" section, read from the
source so the two cannot drift apart.

    python docs/softwarex/build_declaration.py   # -> docs/softwarex/declaration_of_interest.docx
"""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "softwarex.tex"
OUT = HERE / "declaration_of_interest.docx"

NONE = ("The authors declare that they have no known competing financial interests "
        "or personal relationships that could have appeared to influence the work "
        "reported in this paper.")
SOME = ("The authors declare the following financial interests/personal relationships "
        "which may be considered as potential competing interests:")


def manuscript():
    tex = SOURCE.read_text()
    title = re.search(r"\\title\{([^}]*)\}", tex).group(1)
    authors = re.findall(r"\\author\[[^\]]*\]\{([^}\\]*)", tex)
    interest = re.search(r"\\section\*\{Declaration of competing interest\}\s*(.*?)\n\s*\n",
                         tex, re.S).group(1).strip()
    return title, authors, interest.replace("~", " ")


def main():
    title, authors, interest = manuscript()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("Declaration of interests", level=1)
    p = doc.add_paragraph()
    p.add_run("Manuscript: ").bold = True
    p.add_run(title)
    p = doc.add_paragraph()
    p.add_run("Authors: ").bold = True
    p.add_run(", ".join(authors))

    doc.add_paragraph("\u2610 " + NONE)
    doc.add_paragraph("\u2612 " + SOME)
    doc.add_paragraph(interest)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
